import docx
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn, nsdecls
import sys
import os
import re

def set_font(run, font_name='宋体', size_pt=11, bold=False, color_rgb=None):
    run.font.name = font_name
    run.font.size = Pt(size_pt)
    run.bold = bold
    if color_rgb:
        run.font.color.rgb = color_rgb
    
    # 关键：设置中文字体东亚语言属性 w:eastAsia
    rPr = run._r.get_or_add_rPr()
    rFonts = rPr.xpath('w:rFonts')
    if not rFonts:
        rFonts_elem = OxmlElement('w:rFonts')
        rFonts_elem.set(qn('w:ascii'), font_name)
        rFonts_elem.set(qn('w:hAnsi'), font_name)
        rFonts_elem.set(qn('w:eastAsia'), font_name)
        rPr.append(rFonts_elem)
    else:
        rFonts[0].set(qn('w:ascii'), font_name)
        rFonts[0].set(qn('w:hAnsi'), font_name)
        rFonts[0].set(qn('w:eastAsia'), font_name)

def set_cell_shading(tcPr, fill):
    """P2-2 幂等：先清除单元格已有的 w:shd 阴影节点，再追加新的。

    同一 docx 被美化脚本执行多次时，每格最多保留一个阴影节点，
    连续执行两次的文档结构保持不变。"""
    for old in tcPr.findall(qn("w:shd")):
        tcPr.remove(old)
    tcPr.append(parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill}"/>'))


def replace_in_runs(runs, old_str, new_str):
    """跨 Run 边界精确替换文本，保留未被替换字符的样式与格式。"""
    if not runs:
        return
    while True:
        full_text = ''
        offsets = []
        for r in runs:
            start = len(full_text)
            full_text += r.text or ''
            offsets.append((start, len(full_text)))
        
        pos = full_text.find(old_str)
        if pos == -1:
            break
        end_pos = pos + len(old_str)
        
        start_r = end_r = None
        for i, (s, e) in enumerate(offsets):
            if s <= pos < e:
                start_r = i
            if s < end_pos <= e:
                end_r = i
        
        if start_r is None or end_r is None:
            for r in runs:
                if old_str in (r.text or ''):
                    r.text = r.text.replace(old_str, new_str)
            break
            
        if start_r == end_r:
            r = runs[start_r]
            s, e = offsets[start_r]
            local_pos = pos - s
            r.text = r.text[:local_pos] + new_str + r.text[local_pos + len(old_str):]
        else:
            s_run = runs[start_r]
            e_run = runs[end_r]
            s_offset, _ = offsets[start_r]
            e_offset, _ = offsets[end_r]
            
            prefix = (s_run.text or '')[:pos - s_offset]
            suffix = (e_run.text or '')[end_pos - e_offset:]
            
            s_run.text = prefix + new_str
            for mid in range(start_r + 1, end_r):
                runs[mid].text = ''
            e_run.text = suffix

def replace_checkboxes(p):
    """替换段落中各种未渲染的方框标记为标准中文字符 □ / ☑，彻底化解跨 Run 节点静默失效问题"""
    if '[ ]' in p.text:
        replace_in_runs(p.runs, '[ ]', '□')
    if '[x]' in p.text:
        replace_in_runs(p.runs, '[x]', '☑')
    if '[X]' in p.text:
        replace_in_runs(p.runs, '[X]', '☑')
    if '☐' in p.text:
        replace_in_runs(p.runs, '☐', '□')

def beautify_lesson_plan_docx(docx_path):
    doc = docx.Document(docx_path)
    
    # 1. 页面设置：标准 A4，优雅边距 (上下 20mm, 左右 22mm)
    for section in doc.sections:
        section.page_width = Inches(8.27)   # 210 mm
        section.page_height = Inches(11.69) # 297 mm
        section.top_margin = Inches(0.79)   # 20 mm
        section.bottom_margin = Inches(0.79)
        section.left_margin = Inches(0.87)  # 22 mm
        section.right_margin = Inches(0.87)

    # 2. 遍历处理段落
    for p in doc.paragraphs:
        replace_checkboxes(p)
        text = p.text.strip()
        if not text:
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(2)
            continue
        
        style_name = p.style.name.lower()
        
        # 顶级标题 (Title / Heading 1)
        if 'heading 1' in style_name or 'title' in style_name:
            p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.first_line_indent = Pt(0)
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(10)
            p.paragraph_format.line_spacing = 1.3
            for run in p.runs:
                set_font(run, font_name='黑体', size_pt=18, bold=True, color_rgb=RGBColor(0x2E, 0x1F, 0x5E))
        
        # 二级标题 (Heading 2，例如 一、教材分析、二、学情分析 等)
        elif 'heading 2' in style_name:
            p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.first_line_indent = Pt(0) # 标题首行不缩进
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.line_spacing = 1.3
            for run in p.runs:
                set_font(run, font_name='黑体', size_pt=13, bold=True, color_rgb=RGBColor(0x2E, 0x1F, 0x5E))
                
        # 三级标题 (Heading 3，例如 ### 1. 导学单 等)
        elif 'heading 3' in style_name:
            p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.first_line_indent = Pt(0)
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.line_spacing = 1.3
            for run in p.runs:
                set_font(run, font_name='黑体', size_pt=11.5, bold=True, color_rgb=RGBColor(0x40, 0x30, 0x70))
        
        # 列表项 (List Bullet / Compact List)
        elif 'compact' in style_name or 'list' in style_name or text.startswith('•') or text.startswith('- '):
            p.paragraph_format.first_line_indent = Pt(0) # 列表项保持悬挂，不加首行缩进
            p.paragraph_format.left_indent = Pt(14)
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.line_spacing = 1.3
            for run in p.runs:
                set_font(run, font_name='宋体', size_pt=10.5, bold=run.bold)
        
        # 普通正文自然段 (First Paragraph, Body Text, Normal)
        else:
            # 关键：正文自然段首行空两格 (2 * 11pt = 22pt)
            p.paragraph_format.first_line_indent = Pt(22)
            p.paragraph_format.left_indent = Pt(0)
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.line_spacing = 1.38 # 舒适的中文行距
            for run in p.runs:
                set_font(run, font_name='宋体', size_pt=11, bold=run.bold)

    # 3. 处理表格：赋予正统中文教案网格全边框，消除断层
    for t_idx, table in enumerate(doc.tables):
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        tblPr = table._tbl.tblPr
        
        # 移除旧边框并设置统一的浅灰细实线全边框 (Table Grid)
        for b in tblPr.xpath('w:tblBorders'):
            tblPr.remove(b)
        borders_xml = f'''
        <w:tblBorders {nsdecls("w")}>
            <w:top w:val="single" w:sz="4" w:space="0" w:color="B0B0B0"/>
            <w:left w:val="single" w:sz="4" w:space="0" w:color="B0B0B0"/>
            <w:bottom w:val="single" w:sz="4" w:space="0" w:color="B0B0B0"/>
            <w:right w:val="single" w:sz="4" w:space="0" w:color="B0B0B0"/>
            <w:insideH w:val="single" w:sz="4" w:space="0" w:color="C8C8C8"/>
            <w:insideV w:val="single" w:sz="4" w:space="0" w:color="C8C8C8"/>
        </w:tblBorders>
        '''
        tblPr.append(parse_xml(borders_xml))
        
        # 单元格舒适内边距 (上下 120 dxa ≈ 6pt, 左右 160 dxa ≈ 8pt)
        for m in tblPr.xpath('w:tblCellMar'):
            tblPr.remove(m)
        cell_mar_xml = f'''
        <w:tblCellMar {nsdecls("w")}>
            <w:top w:w="120" w:type="dxa"/>
            <w:bottom w:w="120" w:type="dxa"/>
            <w:left w:w="160" w:type="dxa"/>
            <w:right w:w="160" w:type="dxa"/>
        </w:tblCellMar>
        '''
        tblPr.append(parse_xml(cell_mar_xml))

        # 教学过程表 (Table 1) 4 列黄金宽度比保护：彻底防止首列“教学环节”折行双排
        # 教学环节 1650 dxa (~29mm), 教师活动 4100 dxa (~72mm), 学生活动 2350 dxa (~41mm), 设计意图 1312 dxa (~23mm)
        if t_idx == 1 and len(table.columns) == 4:
            col_widths = [1650, 4100, 2350, 1312]
            tblGrid = table._tbl.xpath('w:tblGrid')
            if tblGrid:
                table._tbl.remove(tblGrid[0])
            grid_xml = f'<w:tblGrid {nsdecls("w")}>'
            for w in col_widths:
                grid_xml += f'<w:gridCol w:w="{w}"/>'
            grid_xml += '</w:tblGrid>'
            table._tbl.insert(0, parse_xml(grid_xml))
            for row in table.rows:
                for c_idx, w in enumerate(col_widths):
                    tcPr = row.cells[c_idx]._tc.get_or_add_tcPr()
                    for old_w in tcPr.xpath('w:tcW'):
                        tcPr.remove(old_w)
                    tcPr.append(parse_xml(f'<w:tcW {nsdecls("w")} w:w="{w}" w:type="dxa"/>'))

        # 给所有单元格设定字体与背景
        for r_idx, row in enumerate(table.rows):
            # 表头行设置跨页自动重复
            if r_idx == 0:
                trPr = row._tr.get_or_add_trPr()
                if not trPr.xpath('w:tblHeader'):
                    trPr.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))
                
            for c_idx, cell in enumerate(row.cells):
                cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                tcPr = cell._tc.get_or_add_tcPr()
                
                # 顶部基本信息表（Table 0）美化：标签列加浅灰底色
                if t_idx == 0:
                    if c_idx in [0, 2]:
                        set_cell_shading(tcPr, "F2F4F7")
                
                # 教学过程表（Table 1）：表头加淡蓝灰底色，各环节首列居中
                elif t_idx == 1:
                    if r_idx == 0:
                        set_cell_shading(tcPr, "E9EDF4")
                    elif c_idx == 0:
                        set_cell_shading(tcPr, "FAFAFC")
                
                # 处理单元格内的段落：表格内首行绝不缩进！
                for p in cell.paragraphs:
                    replace_checkboxes(p)
                    p.paragraph_format.first_line_indent = Pt(0)
                    p.paragraph_format.space_before = Pt(1)
                    p.paragraph_format.space_after = Pt(1)
                    p.paragraph_format.line_spacing = 1.25
                    
                    is_header = (t_idx == 0 and c_idx in [0, 2]) or (t_idx >= 1 and r_idx == 0)

                    # 水平对齐：表头全体居中、教学过程表“教学环节”首列居中、基本信息表标签列居中；其余内容靠左
                    if is_header or (t_idx == 1 and c_idx == 0):
                        p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    else:
                        p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT

                    for run in p.runs:
                        font_name = '黑体' if is_header else '宋体'
                        size = 10 if t_idx == 1 else 10.5
                        set_font(run, font_name=font_name, size_pt=size, bold=(is_header or run.bold))

    doc.save(docx_path)
    print(f"SUCCESS: Beautified lesson plan docx: {docx_path}")

def beautify_worksheet_docx(docx_path):
    doc = docx.Document(docx_path)

    # 导学单页面设置：统一采用上下 14mm (0.55 in)、左右 17mm (0.67 in)，
    # 留足纵向安全裕度，严禁因多出 1~2 行溢出（单页）或超页（双页）。
    # 注：此前曾有"单页自适应放大间距"的计划，但两个分支参数完全相同且
    # 后续排版未引用该判定，已于 2026-10-02 清理为统一边距，避免误导。
    for section in doc.sections:
        section.page_width = Inches(8.27)   # 210 mm
        section.page_height = Inches(11.69) # 297 mm
        section.top_margin = Inches(0.55)
        section.bottom_margin = Inches(0.55)
        section.left_margin = Inches(0.67)
        section.right_margin = Inches(0.67)

    for p in doc.paragraphs:
        replace_checkboxes(p)
        text = p.text.strip()
        if not text:
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(1)
            continue
            
        style_name = p.style.name.lower()
        
        # 导学单主标题（段后留出整整一行的舒展留白，告别“脖子太短”）
        if 'heading 1' in style_name or 'title' in style_name:
            p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.first_line_indent = Pt(0)
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(18)
            p.paragraph_format.line_spacing = 1.2
            for run in p.runs:
                set_font(run, font_name='黑体', size_pt=16, bold=True, color_rgb=RGBColor(0x1F, 0x2A, 0x44))
        
        # 班级姓名信息行（右对齐，并在标题与正文间留足舒展空隙，自动剥离重复的今日评价）
        elif '班级' in text and '姓名' in text and len(text) < 60:
            p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            p.paragraph_format.first_line_indent = Pt(0)
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(14)
            p.paragraph_format.line_spacing = 1.15
            
            # 若仍残留今日评价/星级，自动清洗剔除（因底部已有自评，避免重复）
            if '评价' in p.text:
                for target in ['今日评价：☆☆☆☆☆', '今日评价：⭐⭐⭐⭐⭐', '今日评价：☆☆☆', '今日评价：⭐⭐⭐', '评价：☆☆☆☆☆', '评价：⭐⭐⭐⭐⭐', '评价：☆☆☆', '评价：⭐⭐⭐', '今日评价', '评价']:
                    if target in p.text:
                        replace_in_runs(p.runs, target, '')
                for r in p.runs:
                    r.text = re.sub(r'[\s:：&nbsp;☆★⭐·]+$', '', r.text)

            for run in p.runs:
                set_font(run, font_name='楷体', size_pt=10, color_rgb=RGBColor(0x33, 0x33, 0x33))
                
        # 通关目标 / 引用块
        elif 'quote' in style_name or 'block text' in style_name or text.startswith('>') or '通关目标' in text:
            p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.first_line_indent = Pt(0)
            p.paragraph_format.left_indent = Pt(6)
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(8)
            p.paragraph_format.line_spacing = 1.2
            for run in p.runs:
                set_font(run, font_name='楷体', size_pt=9.5, bold=run.bold, color_rgb=RGBColor(0x2E, 0x40, 0x5E))
                
        # 关卡标题 (Heading 3，例如 ### ⭐ 第一关...)
        elif 'heading 3' in style_name or 'heading 2' in style_name or text.startswith('###') or '关：' in text or '打卡' in text:
            p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.first_line_indent = Pt(0)
            p.paragraph_format.space_before = Pt(9) # 关卡间适度留白，不占多余高度
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.line_spacing = 1.2
            for run in p.runs:
                set_font(run, font_name='黑体', size_pt=11.5, bold=True, color_rgb=RGBColor(0x1F, 0x3A, 0x60))

        # 自评总结行 (例如: 今日我的自评：□ ⭐ 基础通关 ...)
        elif '今日我的自评' in text or '挑战自评' in text or ('自评' in text and ('⭐' in text or '□' in text)):
            pPr = p._p.get_or_add_pPr()
            for numPr in pPr.xpath('w:numPr'):
                pPr.remove(numPr)
            p.paragraph_format.first_line_indent = Pt(0)
            p.paragraph_format.left_indent = Pt(0)
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.15
            for run in p.runs:
                set_font(run, font_name='宋体', size_pt=10, bold=run.bold)
                
        # 列表条目 / 选项勾选项 / 序号步骤 (保持悬挂对齐，首行不缩进)
        elif (
            text.startswith('□') or text.startswith('☑') or text.startswith('- □') 
            or text.startswith('- [ ]') or ('□' in text and text.startswith('-'))
            or text.startswith('- ') or text.startswith('•')
            or 'compact' in style_name
            or bool(re.match(r'^[0-9]+[\.、\)]', text))
            or bool(re.match(r'^[A-Za-z][\.、\)]', text))
            or bool(re.match(r'^[①-⑩]', text))
        ):
            pPr = p._p.get_or_add_pPr()
            for numPr in pPr.xpath('w:numPr'):
                pPr.remove(numPr)
            p.paragraph_format.first_line_indent = Pt(0) # 列表项保持左侧悬挂对齐
            p.paragraph_format.left_indent = Pt(14)
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.2
            for run in p.runs:
                set_font(run, font_name='宋体', size_pt=10, bold=run.bold)
                
        # 核心：普通题干与引导正文段落（用户诉求 1：导学案需要缩进）
        else:
            # 正文自然段/题干说明首行空两格 (2 * 10.5pt = 21pt)
            p.paragraph_format.first_line_indent = Pt(21)
            p.paragraph_format.left_indent = Pt(0)
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.line_spacing = 1.22
            for run in p.runs:
                set_font(run, font_name='宋体', size_pt=10.5, bold=run.bold)

    # 处理导学单内部表格（如连线匹配表）
    for table in doc.tables:
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        tblPr = table._tbl.tblPr
        
        for b in tblPr.xpath('w:tblBorders'):
            tblPr.remove(b)
        borders_xml = f'''
        <w:tblBorders {nsdecls("w")}>
            <w:top w:val="single" w:sz="4" w:space="0" w:color="A0A0A0"/>
            <w:left w:val="single" w:sz="4" w:space="0" w:color="A0A0A0"/>
            <w:bottom w:val="single" w:sz="4" w:space="0" w:color="A0A0A0"/>
            <w:right w:val="single" w:sz="4" w:space="0" w:color="A0A0A0"/>
            <w:insideH w:val="single" w:sz="4" w:space="0" w:color="C0C0C0"/>
            <w:insideV w:val="single" w:sz="4" w:space="0" w:color="C0C0C0"/>
        </w:tblBorders>
        '''
        tblPr.append(parse_xml(borders_xml))
        
        # 单元格内边距：适度舒适，上下 75 dxa (约3.8pt)，左右 120 dxa
        cell_top = "75"
        cell_side = "120"
        for m in tblPr.xpath('w:tblCellMar'):
            tblPr.remove(m)
        cell_mar_xml = f'''
        <w:tblCellMar {nsdecls("w")}>
            <w:top w:w="{cell_top}" w:type="dxa"/>
            <w:bottom w:w="{cell_top}" w:type="dxa"/>
            <w:left w:w="{cell_side}" w:type="dxa"/>
            <w:right w:w="{cell_side}" w:type="dxa"/>
        </w:tblCellMar>
        '''
        tblPr.append(parse_xml(cell_mar_xml))
        
        for r_idx, row in enumerate(table.rows):
            for c_idx, cell in enumerate(row.cells):
                cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                tcPr = cell._tc.get_or_add_tcPr()
                if r_idx == 0:
                    set_cell_shading(tcPr, "EDF2F7")
                
                for p in cell.paragraphs:
                    replace_checkboxes(p)
                    p.paragraph_format.first_line_indent = Pt(0)
                    p.paragraph_format.space_before = Pt(1)
                    p.paragraph_format.space_after = Pt(1)
                    p.paragraph_format.line_spacing = 1.15
                    if r_idx == 0:
                        p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    for run in p.runs:
                        is_h = (r_idx == 0)
                        size_pt = 9.5
                        set_font(run, font_name='黑体' if is_h else '宋体', size_pt=size_pt, bold=is_h or run.bold)

    doc.save(docx_path)
    print(f"SUCCESS: Beautified worksheet docx: {docx_path}")

def auto_beautify_docx(docx_path):
    base = os.path.basename(docx_path)
    if '导学' in base or '学习' in base or '任务单' in base or '作业' in base:
        beautify_worksheet_docx(docx_path)
    else:
        beautify_lesson_plan_docx(docx_path)

if __name__ == '__main__':
    target = sys.argv[1] if len(sys.argv) > 1 else 'examples/三年级上/第02课_了解智能工具/02_了解智能工具_教案.docx'
    auto_beautify_docx(target)
